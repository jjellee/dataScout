#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ir_events.py - 기업설명회(IR) 개최 공시 → 누적 엑셀 + 구글 캘린더(ICS 구독)

DART에 이미 받아 둔 「기업설명회(IR)개최」·「기업설명회(IR)개최(안내공시)」 원문을 파싱한다.
서식이 두 가지다(정정공시는 정정사항 표 뒤에 같은 서식이 다시 붙는다).
  A. 안내공시(주로 유가)  1. 일시 및 장소 | 일시 | 날짜 | 시간 | 장소 | … 7. 결정일자 | 8. IR 자료
  B. 개최(주로 코스닥)    1. 일시 | 행사일 | 시작일 | 종료일 | 시작시간 | 종료시간 | 2. 장소 | …

  (기본)/--daily     신규 원문 파싱 → 엑셀·ICS → 텔레그램 (평일 21시 run_dart_periodic.sh)
  --parse            저장된 원문 전량 재파싱
  --backfill A B     기간 원문 보충(직렬) + 파싱
  --excel            캐시로 엑셀·ICS만 생성
  --no-send / --test

캐시: data_dart/ir_events_cache.json
산출: data_dart/ir_events.xlsx  (예정 IR / 증권사 행사별 / 종목별 / 원장)
      data_dart/ir_calendar.ics       — 전 종목
      data_dart/ir_calendar_major.ics — 시총 1조↑ + watchlist.txt
ICS는 run_daily.sh가 GitHub(공개 저장소)에 커밋하므로 구글 캘린더 「URL로 추가」로 구독한다.
  https://raw.githubusercontent.com/jjellee/dataScout/main/data_dart/ir_calendar_major.ics
"""
import os
import re
import sys
import glob
import html
import json
import argparse
import datetime
import hashlib
import logging

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s',
                    handlers=[logging.StreamHandler(sys.stdout)])
logger = logging.getLogger("ir_events")

WORKSPACE = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(WORKSPACE, "data_dart")
CACHE_PATH = os.path.join(DATA_DIR, "ir_events_cache.json")
EXCEL_PATH = os.path.join(DATA_DIR, "ir_events.xlsx")
ICS_ALL = os.path.join(DATA_DIR, "ir_calendar.ics")
ICS_MAJOR = os.path.join(DATA_DIR, "ir_calendar_major.ics")
RAW_URL = "https://raw.githubusercontent.com/jjellee/dataScout/main/data_dart/"

MAJOR_MCAP = 1_0000_0000_0000      # 1조
ICS_PAST_DAYS = 120                # 캘린더엔 최근 4개월 + 예정분만 싣는다

_env = os.path.join(WORKSPACE, ".env")
if os.path.exists(_env):
    with open(_env, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, v = line.split("=", 1)
                os.environ.setdefault(k.strip(), v.strip().strip("'\""))

sys.path.insert(0, WORKSPACE)


# ------------------------------------------------------------------ 파싱 --- #
DATE_RE = re.compile(r"(20\d\d)[-./년\s]+(\d{1,2})[-./월\s]+(\d{1,2})")
TIME_RE = re.compile(r"(\d{1,2})\s*[:시]\s*(\d{2})?")

# 시간(현지시간)이 한국 시간이 아닐 수 있는 해외 행사 — 종일 일정으로 싣는다
OVERSEAS = ("미국", "뉴욕", "보스턴", "샌프란시스코", "LA", "라스베가스", "라스베이거스",
            "홍콩", "싱가포르", "싱가폴", "일본", "도쿄", "동경", "런던", "영국", "유럽",
            "프랑크푸르트", "파리", "독일", "중국", "상해", "상하이", "북경", "베이징", "대만",
            "타이베이", "호주", "시드니", "캐나다", "토론토", "두바이", "인도", "해외",
            "New York", "Hong Kong", "Singapore", "London", "Tokyo", "Boston",
            "San Francisco", "Las Vegas")


def html_to_cells(path):
    """원문 HTML → 셀 텍스트 리스트 (표 구조 대신 순서만 쓴다)."""
    with open(path, encoding="utf-8", errors="ignore") as f:
        t = f.read()
    t = re.sub(r"(?is)<(style|script)[^>]*>.*?</\1>", " ", t)
    t = re.sub(r"<[^>]+>", "|", t)
    t = html.unescape(t).replace("\xa0", " ")
    cells = [re.sub(r"\s+", " ", c).strip() for c in t.split("|")]
    return [c for c in cells if c]


def _date(s):
    m = DATE_RE.search(s or "")
    if not m:
        return None
    try:
        return datetime.date(int(m[1]), int(m[2]), int(m[3])).isoformat()
    except ValueError:
        return None


def _time(s):
    m = TIME_RE.search(s or "")
    if not m:
        return None
    h, mi = int(m[1]), int(m[2] or 0)
    if h > 23 or mi > 59:
        return None
    return f"{h:02d}:{mi:02d}"


def _between(cells, start_pat, end_pats):
    """start_pat에 맞는 셀 다음부터 end_pats 중 하나가 나올 때까지의 셀."""
    for i, c in enumerate(cells):
        if re.match(start_pat, c):
            out = []
            for d in cells[i + 1:]:
                if any(re.match(p, d) for p in end_pats):
                    break
                out.append(d)
            return out
    return []


SEC = r"\d{1,2}\s*\.\s*"


def _sec(cells, label, nxt=None):
    """'N. 라벨' 다음부터 다음 'N. ' 섹션 전까지."""
    end = [rf"^{SEC}\S", r"^※"]
    return _between(cells, rf"^{SEC}{label}", end)


def parse_ir_html(path):
    cells = html_to_cells(path)
    rec = {}

    # 정정공시: 정정사항 표 뒤에 정정 후 서식 전체가 다시 붙는다 → 마지막 서식만 본다
    is_amend = any(c.startswith("정정신고") or "정정관련 공시서류" in c for c in cells)
    if is_amend:
        rec["정정"] = True
        orig = _sec(cells, r"정정관련 공시서류제출일")
        rec["원공시일"] = _date(" ".join(orig))
        rec["정정사유"] = " ".join(_sec(cells, r"정정사유"))[:200]
        rec["취소"] = "취소" in rec["정정사유"] or "철회" in rec["정정사유"]
        # 정정 전 행사일 (중복 원공시를 고를 때 쓴다)
        tbl = _sec(cells, r"정정사항")
        for i, c in enumerate(tbl):
            if ("일시" in c or "행사일" in c or "시작일" in c) and "게재" not in c and i + 1 < len(tbl):
                rec["정정전_행사일"] = _date(tbl[i + 1])
                break
        # 서식 본문 시작점(마지막 '1. 일시')부터 자른다
        idx = [i for i, c in enumerate(cells) if re.match(rf"^{SEC}일시", c)]
        if idx:
            cells = cells[idx[-1]:]

    def j(lst):
        return " ".join(lst).strip(" -") or ""

    if any(re.match(rf"^{SEC}일시 및 장소", c) for c in cells):       # 서식 A
        rec["form"] = "A"
        blk = _sec(cells, r"일시 및 장소")
        # 일시 | 날짜 | 시간 | 장소 | 값
        try:
            i = blk.index("일시")
        except ValueError:
            i = -1
        try:
            k = blk.index("장소")
        except ValueError:
            k = len(blk)
        when = blk[i + 1:k] if i >= 0 else blk[:k]
        dates = [d for d in (_date(x) for x in when) if d]
        rec["시작일"] = dates[0] if dates else None
        rec["종료일"] = dates[-1] if dates else None
        tpart = [x for x in when if not _date(x)]
        times = re.findall(r"\d{1,2}\s*:\s*\d{2}", " ".join(tpart))
        rec["시작시간"] = _time(times[0]) if times else None
        rec["종료시간"] = _time(times[1]) if len(times) > 1 else None
        rec["시간원문"] = j(tpart)
        rec["장소"] = j(blk[k + 1:])
        rec["대상"] = j(_sec(cells, r"참가\s*대상"))
        rec["목적"] = j(_sec(cells, r"개최\s*목적"))
        rec["방법"] = j(_sec(cells, r"개최\s*방법"))
        rec["주관"] = j(_sec(cells, r"후원\s*기관"))
        rec["내용"] = j(_sec(cells, r"주요\s*설명회"))
        rec["결정일"] = _date(j(_sec(cells, r"결정일자")))
        ir = _sec(cells, r"IR\s*자료")
        rec["자료게재"] = _date(j(ir))
        web = [x for x in ir if re.search(r"https?://|www\.|\.co\.kr|\.com", x)]
        rec["웹페이지"] = web[0] if web else ""
        rec["기타"] = j(_sec(cells, r"기타\s*투자판단"))[:500]
    else:                                                              # 서식 B
        rec["form"] = "B"
        blk = _sec(cells, r"일시")
        vals = [x for x in blk if x not in ("행사일", "시간(현지시간)", "시작일", "종료일",
                                            "시작시간", "종료시간", "시간")]
        dates = [d for d in (_date(x) for x in vals) if d]
        rest = [x for x in vals if not _date(x)]
        rec["시작일"] = dates[0] if dates else None
        rec["종료일"] = dates[1] if len(dates) > 1 else rec["시작일"]
        rec["시작시간"] = _time(rest[0]) if rest else None
        rec["종료시간"] = _time(rest[1]) if len(rest) > 1 else None
        rec["시간원문"] = j(rest)
        rec["장소"] = j(_sec(cells, r"장소"))
        rec["대상"] = j(_sec(cells, r"대상자"))
        rec["목적"] = j(_sec(cells, r"실시\s*목적"))
        rec["방법"] = j(_sec(cells, r"실시\s*방법"))
        rec["주관"] = j(_sec(cells, r"후원\s*기관"))
        rec["내용"] = j(_sec(cells, r"주요\s*내용"))
        rec["결정일"] = _date(j(_sec(cells, r"개최\s*확정일")))
        ir = _sec(cells, r"IR\s*자료\s*게재")
        rec["자료게재"] = _date(j(ir))
        web = [x for x in ir if re.search(r"https?://|www\.|\.co\.kr|\.com", x)]
        rec["웹페이지"] = web[0] if web else ""
        담당 = _sec(cells, r"담당자")
        rec["담당"] = j([x for x in 담당 if x not in ("담당부서(담당자)", "전화번호(팩스번호)")])
        rec["기타"] = ""

    if rec.get("종료일") and rec.get("시작일") and rec["종료일"] < rec["시작일"]:
        rec["종료일"] = rec["시작일"]
    rec["구분"] = classify(rec)
    rec["해외"] = is_overseas(rec)
    return rec


def is_overseas(r):
    s = f"{r.get('장소', '')} {r.get('방법', '')}"
    if re.search(r"국내|서울|여의도|부산|판교", r.get("장소", "")):
        return False
    return any(k in s for k in OVERSEAS)


def classify(r):
    s = " ".join(r.get(k, "") for k in ("목적", "방법", "내용", "장소"))
    if re.search(r"실적|결산|분기|반기|Earnings|earnings|컨퍼런스\s*콜|Conference\s*Call", s) \
            and not re.search(r"Corporate\s*Day|콥\s*데이|NDR|Conference(?!\s*Call)|포럼|박람회", s, re.I):
        return "실적발표"
    if re.search(r"NDR|Non[- ]?Deal", s, re.I):
        return "NDR"
    if re.search(r"Corporate\s*Day|Corp\.?\s*Day|콥\s*데이|코퍼레이트\s*데이|IR\s*Day|Investor\s*Day|기업설명회\s*참[가여]|컨퍼런스|Conference|포럼|Forum|"
                 r"Summit|세미나|증권.*(?:주최|초청|주관)|IR\s*협의회|거래소", s, re.I):
        return "증권사·기관 행사"
    if re.search(r"밸류업|기업가치\s*제고|주주환원", s):
        return "밸류업·주주환원"
    if re.search(r"개인투자자|일반투자자|주주간담|주주\s*소통", s):
        return "개인·주주"
    if re.search(r"1\s*:\s*1|일대일|대면|미팅|탐방", s):
        return "1:1·소규모 미팅"
    return "기타"


# ------------------------------------------------------------- 수집·캐시 --- #
def load_cache():
    if os.path.exists(CACHE_PATH):
        with open(CACHE_PATH, encoding="utf-8") as f:
            return json.load(f)
    return {"events": {}}


def save_cache(cache):
    tmp = CACHE_PATH + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=1)
    os.replace(tmp, CACHE_PATH)


def iter_ir_disclosures(date_from=None, date_to=None):
    for p in sorted(glob.glob(os.path.join(DATA_DIR, "20[0-9][0-9][0-9][0-9][0-9][0-9]", "disclosures.json"))):
        d = os.path.basename(os.path.dirname(p))
        if (date_from and d < date_from) or (date_to and d > date_to):
            continue
        try:
            with open(p, encoding="utf-8") as f:
                items = json.load(f)
        except Exception:
            continue
        for it in items:
            nm = str(it.get("report_nm", "")).replace(" ", "")
            if "기업설명회" not in nm or "결과" in nm:
                continue
            yield os.path.dirname(p), it


def fetch_missing(date_from=None, date_to=None):
    """목록에는 있는데 원문이 없는 IR 공시를 받는다. DART를 병렬로 두드리면 IP가 막히므로 직렬."""
    import time
    from dart_collector import download_disclosure_document
    key = os.getenv("DART_API_KEY")
    if not key:
        logger.warning("DART_API_KEY 없음 — 원문 보충 생략")
        return 0
    todo = [(d, it) for d, it in iter_ir_disclosures(date_from, date_to)
            if not os.path.exists(os.path.join(d, f"{it['rcept_no']}.html"))]
    ok = 0
    for i, (d, it) in enumerate(todo, 1):
        if download_disclosure_document(key, it["rcept_no"], d, metadata=it):
            ok += 1
        if i % 100 == 0:
            logger.info(f"  원문 보충 {i}/{len(todo)}")
        time.sleep(0.3)
    logger.info(f"IR 원문 보충 {ok}/{len(todo)}건")
    return ok


def parse_all(cache, date_from=None, date_to=None, reparse=False):
    ev = cache["events"]
    n_new = n_miss = n_fail = 0
    for ddir, it in iter_ir_disclosures(date_from, date_to):
        rno = it["rcept_no"]
        if rno in ev and not reparse:
            continue
        path = os.path.join(ddir, f"{rno}.html")
        if not os.path.exists(path):
            n_miss += 1
            continue
        try:
            rec = parse_ir_html(path)
        except Exception as e:
            logger.warning(f"{rno} 파싱 실패: {e}")
            n_fail += 1
            continue
        rec.update({
            "rcept_no": rno, "rcept_dt": it.get("rcept_dt"), "corp_name": it.get("corp_name"),
            "stock_code": it.get("stock_code") or "", "corp_cls": it.get("corp_cls"),
            "report_nm": re.sub(r"\s+", " ", it.get("report_nm", "")).strip(),
        })
        if not rec.get("시작일"):
            n_fail += 1
        ev[rno] = rec
        n_new += 1
    logger.info(f"IR 공시 파싱 {n_new}건 (원문 없음 {n_miss} — 2025년 이전 미수집분 포함, 일자 미파싱 {n_fail})")
    return n_new


def resolve_events(cache):
    """정정·중복을 걸러 실제 행사 목록을 만든다.

    - 정정공시는 같은 회사의 원공시(원공시일 접수분)를 대체한다. 원공시가 여럿이면
      정정 전 행사일이 같은 것만 대체한다.
    - 같은 회사·같은 날·같은 시작시간·같은 장소는 한 행사로 보고 최신 접수분만 남긴다
      (국문·영문 동시 공시, 안내공시 재공시).
    """
    allr = list(cache["events"].values())
    superseded = set()
    by_corp_dt = {}
    for r in allr:
        by_corp_dt.setdefault((r["stock_code"] or r["corp_name"], r["rcept_dt"]), []).append(r)
    # 접수 순서대로 — 정정의 정정도 앞 정정을 대체한다
    for r in sorted(allr, key=lambda x: x["rcept_no"]):
        if not r.get("정정") or not r.get("원공시일"):
            continue
        cands = [c for c in by_corp_dt.get((r["stock_code"] or r["corp_name"],
                                            r["원공시일"].replace("-", "")), [])
                 if c["rcept_no"] < r["rcept_no"]]
        if len(cands) > 1 and r.get("정정전_행사일"):
            cands = [c for c in cands if c.get("시작일") == r["정정전_행사일"]]
        for c in cands:
            superseded.add(c["rcept_no"])
        if r.get("취소"):
            superseded.add(r["rcept_no"])
    recs = [r for r in allr if r.get("시작일")]

    best = {}
    for r in recs:
        if r["rcept_no"] in superseded:
            continue
        key = (r["stock_code"] or r["corp_name"], r["시작일"], r.get("시작시간") or "",
               re.sub(r"\W", "", r.get("장소", ""))[:20])
        if key not in best or r["rcept_no"] > best[key]["rcept_no"]:
            best[key] = r
    return sorted(best.values(), key=lambda r: (r["시작일"], r.get("시작시간") or "99", r["corp_name"]))


# ---------------------------------------------------------------- 보조데이터 --- #
def load_mcap():
    files = sorted(glob.glob(os.path.join(WORKSPACE, "data_kr", "*", "market_cap.csv")))
    if not files:
        return {}
    import csv
    out = {}
    with open(files[-1], encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            try:
                out[row["티커"].zfill(6)] = float(row["시가총액"])
            except (KeyError, ValueError):
                pass
    return out


def load_watchlist():
    p = os.path.join(WORKSPACE, "watchlist.txt")
    if not os.path.exists(p):
        return set()
    with open(p, encoding="utf-8") as f:
        return {ln.strip()[:6] for ln in f if re.match(r"^\d{6}", ln.strip())}


MARKET = {"Y": "유가", "K": "코스닥", "N": "코넥스", "E": "기타"}
WD = "월화수목금토일"


def dart_url(rno):
    return f"https://dart.fss.or.kr/dsaf001/main.do?rcpNo={rno}"


# ------------------------------------------------------------------- 엑셀 --- #
def build_excel(events, out_path, today=None):
    import openpyxl
    from openpyxl.styles import PatternFill, Font
    from quarterly_earnings import _style_sheet, FMT_INT, FONT_NAME

    today = today or datetime.date.today().isoformat()
    mcap = load_mcap()
    watch = load_watchlist()

    def mc(r):
        v = mcap.get(r["stock_code"])
        return round(v / 1e8) if v else None

    def when(r):
        t = r.get("시작시간") or ""
        if r.get("종료시간"):
            t += f"~{r['종료시간']}"
        return t or (r.get("시간원문") or "")[:20]

    def wd(d):
        return WD[datetime.date.fromisoformat(d).weekday()]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    hl_fill = PatternFill("solid", start_color="FFF4D6", end_color="FFF4D6")

    # ── 1) 예정 IR — 오늘 이후, 일자순
    cols = ["행사일", "요일", "시간", "종료일", "종목명", "종목코드", "시장", "시총(억)", "관심",
            "구분", "장소", "주관·후원", "방법", "대상", "목적", "주요내용", "웹페이지",
            "공시일", "정정", "공시 원문"]
    ws = wb.create_sheet("예정 IR")
    ws.append(cols)
    up = [r for r in events if (r.get("종료일") or r["시작일"]) >= today]
    hi_rows = []
    for r in up:
        star = "★" if r["stock_code"] in watch else ""
        ws.append([r["시작일"], wd(r["시작일"]), when(r),
                   r["종료일"] if r.get("종료일") != r["시작일"] else "",
                   r["corp_name"], r["stock_code"], MARKET.get(r.get("corp_cls"), ""), mc(r), star,
                   r["구분"] + (" (해외)" if r.get("해외") else ""), r.get("장소"), r.get("주관"),
                   r.get("방법"), r.get("대상"), r.get("목적"), r.get("내용"), r.get("웹페이지"),
                   r["rcept_dt"], "정정" if r.get("정정") else "", "열기"])
        if star or (mc(r) or 0) >= 100000:
            hi_rows.append(ws.max_row)
        ws.cell(row=ws.max_row, column=len(cols)).hyperlink = dart_url(r["rcept_no"])
    widths = {"행사일": 11, "요일": 5, "시간": 12, "종료일": 11, "종목명": 16, "종목코드": 9,
              "시장": 7, "시총(억)": 10, "관심": 5, "구분": 14, "장소": 26, "주관·후원": 18,
              "방법": 30, "대상": 18, "목적": 30, "주요내용": 36, "웹페이지": 24,
              "공시일": 10, "정정": 6, "공시 원문": 8}
    _style_sheet(ws, headers=cols, freeze="F2", widths=widths,
                 fmts={"시총(억)": FMT_INT}, links=("공시 원문",),
                 center=("요일", "종목코드", "시장", "관심", "정정", "공시 원문"), tab="0F766E")
    for ri in hi_rows:          # 시총 10조↑·관심종목 행 강조
        for ci in range(1, 11):
            ws.cell(row=ri, column=ci).fill = hl_fill
    ws.auto_filter.ref = ws.dimensions

    # ── 2) 증권사 행사별 — 같은 날·같은 주관사에 3사 이상 (Corporate Day·컨퍼런스)
    groups = {}
    for r in events:
        host = re.sub(r"\(.*?\)|주식회사|㈜|\s", "", r.get("주관") or "")
        if not host or host in ("-", "자사", "자사주관", "당사", "없음", "해당없음") or "자사" in host:
            continue
        groups.setdefault((r["시작일"], host[:12]), []).append(r)
    cols2 = ["행사일", "요일", "주관·후원", "참가사 수", "행사명(방법)", "장소", "참가사(시총순)"]
    ws = wb.create_sheet("증권사 행사별")
    ws.append(cols2)
    gl = [(k, v) for k, v in groups.items() if len({x["stock_code"] or x["corp_name"] for x in v}) >= 3]
    # 예정분은 가까운 순, 지난 행사는 최근 순
    fut = sorted([g for g in gl if g[0][0] >= today], key=lambda g: g[0][0])
    past = sorted([g for g in gl if g[0][0] < today], key=lambda g: g[0][0], reverse=True)
    for (d, host), rs in fut + past:
        uniq = {}
        for x in rs:
            uniq.setdefault(x["stock_code"] or x["corp_name"], x)
        rs = sorted(uniq.values(), key=lambda x: -(mcap.get(x["stock_code"]) or 0))
        names = ", ".join(x["corp_name"] + ("★" if x["stock_code"] in watch else "") for x in rs)
        meth = max((x.get("방법") or "" for x in rs), key=lambda s: sum(
            (y.get("방법") or "") == s for y in rs))
        ws.append([d, wd(d), rs[0].get("주관"), len(rs), meth, rs[0].get("장소"), names])
    _style_sheet(ws, headers=cols2, freeze="C2",
                 widths={"행사일": 11, "요일": 5, "주관·후원": 18, "참가사 수": 9,
                         "행사명(방법)": 40, "장소": 26, "참가사(시총순)": 120},
                 fmts={"참가사 수": FMT_INT}, center=("요일", "참가사 수"), tab="1F3B57")
    for ri in range(2, ws.max_row + 1):
        if ws.cell(row=ri, column=1).value >= today:
            for ci in range(1, 5):
                ws.cell(row=ri, column=ci).fill = hl_fill
    ws.auto_filter.ref = ws.dimensions

    # ── 3) 종목별 — 최근 12개월 IR 횟수·다음 일정
    since = (datetime.date.fromisoformat(today) - datetime.timedelta(days=365)).isoformat()
    corp = {}
    for r in events:
        k = r["stock_code"] or r["corp_name"]
        c = corp.setdefault(k, {"r": r, "n12": 0, "earn": 0, "ndr": 0, "conf": 0,
                                "last": None, "next": None, "hosts": {}})
        if since <= r["시작일"] < today:
            c["n12"] += 1
            c["earn"] += r["구분"] == "실적발표"
            c["ndr"] += r["구분"] == "NDR"
            c["conf"] += r["구분"] == "증권사·기관 행사"
            c["last"] = max(c["last"] or "", r["시작일"])
            h = (r.get("주관") or "").strip()
            if h and "자사" not in h and h != "-":
                c["hosts"][h] = c["hosts"].get(h, 0) + 1
        if r["시작일"] >= today and (c["next"] is None or r["시작일"] < c["next"][0]):
            c["next"] = (r["시작일"], r["구분"])
    cols3 = ["종목명", "종목코드", "시장", "시총(억)", "관심", "12개월 IR", "실적발표", "NDR",
             "증권사·기관", "최근 IR", "다음 IR", "다음 IR 구분", "주요 주관사"]
    ws = wb.create_sheet("종목별")
    ws.append(cols3)
    for k, c in sorted(corp.items(), key=lambda kv: (-kv[1]["n12"], kv[1]["r"]["corp_name"])):
        r = c["r"]
        hosts = ", ".join(h for h, _ in sorted(c["hosts"].items(), key=lambda x: -x[1])[:4])
        ws.append([r["corp_name"], r["stock_code"], MARKET.get(r.get("corp_cls"), ""), mc(r),
                   "★" if r["stock_code"] in watch else "", c["n12"], c["earn"], c["ndr"],
                   c["conf"], c["last"] or "", c["next"][0] if c["next"] else "",
                   c["next"][1] if c["next"] else "", hosts])
    _style_sheet(ws, headers=cols3, freeze="C2",
                 widths={"종목명": 18, "종목코드": 9, "시장": 7, "시총(억)": 11, "관심": 5,
                         "12개월 IR": 9, "실적발표": 8, "NDR": 6, "증권사·기관": 9,
                         "최근 IR": 11, "다음 IR": 11, "다음 IR 구분": 13, "주요 주관사": 50},
                 fmts={c: FMT_INT for c in ("시총(억)", "12개월 IR", "실적발표", "NDR", "증권사·기관")},
                 center=("종목코드", "시장", "관심", "12개월 IR", "실적발표", "NDR", "증권사·기관",
                         "최근 IR", "다음 IR"), tab="2E5A80")
    ws.auto_filter.ref = ws.dimensions

    # ── 4) 원장 — 정정·중복 포함 전량, 접수일 역순
    cols4 = ["접수일", "종목명", "종목코드", "시장", "보고서명", "행사일", "종료일", "시간",
             "구분", "장소", "주관·후원", "방법", "대상", "목적", "주요내용", "정정사유",
             "결정일", "자료게재일", "담당", "접수번호"]
    ws = wb.create_sheet("원장")
    ws.append(cols4)
    from_cache = sorted(_ALL_RECS, key=lambda r: r["rcept_no"], reverse=True)
    for r in from_cache:
        ws.append([r["rcept_dt"], r["corp_name"], r["stock_code"], MARKET.get(r.get("corp_cls"), ""),
                   r["report_nm"], r.get("시작일") or "", r.get("종료일") or "", when(r),
                   r.get("구분"), r.get("장소"), r.get("주관"), r.get("방법"), r.get("대상"),
                   r.get("목적"), r.get("내용"), r.get("정정사유", ""), r.get("결정일") or "",
                   r.get("자료게재") or "", r.get("담당", ""), r["rcept_no"]])
        ws.cell(row=ws.max_row, column=len(cols4)).hyperlink = dart_url(r["rcept_no"])
    _style_sheet(ws, headers=cols4, freeze="C2",
                 widths={"접수일": 10, "종목명": 16, "종목코드": 9, "시장": 7, "보고서명": 30,
                         "행사일": 11, "종료일": 11, "시간": 12, "구분": 13, "장소": 24,
                         "주관·후원": 18, "방법": 28, "대상": 18, "목적": 28, "주요내용": 32,
                         "정정사유": 24, "결정일": 11, "자료게재일": 11, "담당": 24, "접수번호": 16},
                 links=("접수번호",), center=("종목코드", "시장"), tab="8A96A3")
    ws.auto_filter.ref = ws.dimensions

    tmp = out_path + ".tmp.xlsx"
    wb.save(tmp)
    os.replace(tmp, out_path)
    logger.info(f"엑셀 저장: {out_path} (예정 {len(up)}건, 행사 {len(fut)}+{len(past)}, 종목 {len(corp)})")
    return out_path, up, fut


_ALL_RECS = []


# -------------------------------------------------------------------- ICS --- #
def _ics_escape(s):
    return (s or "").replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")


def _fold(line):
    """RFC 5545 — 75옥텟 단위로 접는다(UTF-8 글자 경계 유지)."""
    b = line.encode("utf-8")
    if len(b) <= 75:
        return line
    out, cur = [], b""
    for ch in line:
        e = ch.encode("utf-8")
        if len(cur) + len(e) > (75 if not out else 74):
            out.append(cur.decode("utf-8"))
            cur = b""
        cur += e
    out.append(cur.decode("utf-8"))
    return "\r\n ".join(out)


def build_ics(events, path, name, today=None):
    today = today or datetime.date.today().isoformat()
    lo = (datetime.date.fromisoformat(today) - datetime.timedelta(days=ICS_PAST_DAYS)).isoformat()
    stamp = datetime.datetime.now(datetime.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    L = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//dataScout//IR events//KO",
         "CALSCALE:GREGORIAN", "METHOD:PUBLISH", f"X-WR-CALNAME:{name}",
         "X-WR-TIMEZONE:Asia/Seoul", "REFRESH-INTERVAL;VALUE=DURATION:PT6H",
         "X-PUBLISHED-TTL:PT6H",
         "BEGIN:VTIMEZONE", "TZID:Asia/Seoul", "BEGIN:STANDARD", "DTSTART:19700101T000000",
         "TZOFFSETFROM:+0900", "TZOFFSETTO:+0900", "TZNAME:KST", "END:STANDARD", "END:VTIMEZONE"]
    n = 0
    for r in events:
        if (r.get("종료일") or r["시작일"]) < lo:
            continue
        uid = hashlib.md5(f"{r['stock_code'] or r['corp_name']}|{r['시작일']}|{r.get('시작시간')}|"
                          f"{r.get('장소')}".encode()).hexdigest() + "@datascout-ir"
        tag = r["구분"] if r["구분"] not in ("기타",) else "IR"
        title = f"[{tag}] {r['corp_name']}"
        if r.get("주관") and "자사" not in r["주관"] and r["구분"] != "실적발표":
            title += f" · {r['주관'][:15]}"
        desc = "\n".join(x for x in [
            f"종목: {r['corp_name']} ({r['stock_code']})",
            f"시간: {r.get('시작시간') or ''}{('~' + r['종료시간']) if r.get('종료시간') else ''}"
            + (" (현지시간)" if r.get("해외") else ""),
            f"장소: {r.get('장소')}", f"주관·후원: {r.get('주관')}", f"방법: {r.get('방법')}",
            f"대상: {r.get('대상')}", f"목적: {r.get('목적')}", f"내용: {r.get('내용')}",
            f"IR자료: {r.get('웹페이지')}" if r.get("웹페이지") else "",
            f"공시: {dart_url(r['rcept_no'])}" + (" (정정)" if r.get("정정") else ""),
        ] if x and not x.endswith(": "))
        d0 = r["시작일"].replace("-", "")
        d1 = (r.get("종료일") or r["시작일"])
        timed = r.get("시작시간") and not r.get("해외") and d1 == r["시작일"]
        L += ["BEGIN:VEVENT", f"UID:{uid}", f"DTSTAMP:{stamp}"]
        if timed:
            st = r["시작시간"].replace(":", "")
            if r.get("종료시간") and r["종료시간"] > r["시작시간"]:
                en = r["종료시간"].replace(":", "")
            else:
                h, m = map(int, r["시작시간"].split(":"))
                en = f"{min(h + 1, 23):02d}{m:02d}"
            L += [f"DTSTART;TZID=Asia/Seoul:{d0}T{st}00", f"DTEND;TZID=Asia/Seoul:{d0}T{en}00"]
        else:
            end = (datetime.date.fromisoformat(d1) + datetime.timedelta(days=1)).strftime("%Y%m%d")
            L += [f"DTSTART;VALUE=DATE:{d0}", f"DTEND;VALUE=DATE:{end}"]
        L += [f"SUMMARY:{_ics_escape(title)}", f"LOCATION:{_ics_escape(r.get('장소'))}",
              f"DESCRIPTION:{_ics_escape(desc)}", f"URL:{dart_url(r['rcept_no'])}",
              "TRANSP:TRANSPARENT", "END:VEVENT"]
        n += 1
    L.append("END:VCALENDAR")
    with open(path + ".tmp", "w", encoding="utf-8", newline="") as f:
        f.write("\r\n".join(_fold(x) for x in L) + "\r\n")
    os.replace(path + ".tmp", path)
    logger.info(f"ICS 저장: {path} ({n}건)")
    return n


def build_outputs(cache, today=None):
    global _ALL_RECS
    _ALL_RECS = list(cache["events"].values())
    events = resolve_events(cache)
    out, up, fut = build_excel(events, EXCEL_PATH, today=today)
    mcap, watch = load_mcap(), load_watchlist()
    major = [r for r in events
             if (mcap.get(r["stock_code"]) or 0) >= MAJOR_MCAP or r["stock_code"] in watch]
    build_ics(events, ICS_ALL, "IR 일정 (전 종목)", today=today)
    build_ics(major, ICS_MAJOR, "IR 일정 (시총 1조↑·관심)", today=today)
    return out, events, up, fut


# ------------------------------------------------------------------- 메인 --- #
def main():
    ap = argparse.ArgumentParser(description="기업설명회(IR) 공시 → 엑셀 + 캘린더(ICS)")
    ap.add_argument("--daily", action="store_true", help="증분 파싱 + 엑셀·ICS + 발송 (기본)")
    ap.add_argument("--date", default=datetime.date.today().strftime("%Y%m%d"))
    ap.add_argument("--parse", action="store_true", help="저장된 원문 전량 재파싱")
    ap.add_argument("--excel", action="store_true", help="캐시로 엑셀·ICS만 생성")
    ap.add_argument("--backfill", nargs=2, metavar=("FROM", "TO"), help="기간 원문 보충 (YYYYMMDD)")
    ap.add_argument("--no-send", action="store_true")
    ap.add_argument("--test", action="store_true")
    args = ap.parse_args()

    cache = load_cache()
    if args.backfill:
        fetch_missing(*args.backfill)
        parse_all(cache, *args.backfill)
        save_cache(cache)
        build_outputs(cache)
        return
    if args.parse:
        cache = {"events": {}}
        parse_all(cache, reparse=True)
        save_cache(cache)
    elif not args.excel:
        recent = (datetime.date.today() - datetime.timedelta(days=10)).strftime("%Y%m%d")
        fetch_missing(recent)
        parse_all(cache)
        save_cache(cache)

    out, events, up, fut = build_outputs(cache)
    if args.no_send or args.excel or args.parse:
        return

    from dart_report import send_document
    d = args.date
    today_new = [r for r in cache["events"].values() if r.get("rcept_dt") == d]
    mcap, watch = load_mcap(), load_watchlist()
    week_end = (datetime.date.today() + datetime.timedelta(days=7)).isoformat()
    soon = [r for r in up if r["시작일"] <= week_end]
    key = sorted([r for r in soon if (mcap.get(r["stock_code"]) or 0) >= MAJOR_MCAP
                  or r["stock_code"] in watch],
                 key=lambda r: (r["시작일"], -(mcap.get(r["stock_code"]) or 0)))
    lines = [f"🎤 기업설명회(IR) 일정 ({d[:4]}-{d[4:6]}-{d[6:]})",
             f"· 당일 IR 공시 {len(today_new)}건 ㆍ 예정 IR {len(up)}건 (7일 내 {len(soon)}건)",
             f"· 증권사 행사(3사↑) 예정 {len(fut)}건"]
    if key:
        lines.append("· 7일 내 시총 1조↑·관심:")
        for r in key[:12]:
            lines.append(f"   {r['시작일'][5:]} {r['corp_name']} [{r['구분']}]")
        if len(key) > 12:
            lines.append(f"   … 외 {len(key) - 12}건")
    lines.append(f"📅 캘린더 구독: {RAW_URL}ir_calendar_major.ics")
    send_document(out, "\n".join(lines)[:1024], test=args.test)


if __name__ == "__main__":
    main()
